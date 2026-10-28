"""Turns a finished pipeline state into JSON, Markdown, standalone HTML and DOCX."""
from __future__ import annotations

import html
import io
from typing import Any

from processlens.graph import State
from processlens.llm import TECHNIQUE_LABEL
from processlens.stages.score import QUESTIONS

QUADRANT_LABEL = {"quick_win": "Quick win", "strategic": "Strategic", "fill_in": "Fill-in", "reconsider": "Reconsider"}


def to_dict(state: State) -> dict[str, Any]:
    process, doc = state["process"], state["doc"]
    scores = {s.candidate_id: s for s in state.get("scores", [])}
    rows = []
    for c in state.get("candidates", []):
        s = scores.get(c.id)
        rows.append({**c.model_dump(), "score": s.model_dump() if s else None})
    rows.sort(key=lambda r: -((r["score"] or {}).get("impact", 0) + (r["score"] or {}).get("feasibility", 0)))
    return {
        "process": {"id": process.id, "name": process.name, "lanes": process.lanes, "tasks": len(process.tasks),
                    "manual_tasks": len(process.manual_tasks), "handoffs": process.handoffs},
        "document": {"title": doc.title, "sentences": len(doc.sentences), "facts": doc.facts.model_dump()},
        "proposed": len(state.get("proposed", [])),
        "candidates": rows,
        "issues": [i.model_dump() for i in state.get("issues", [])],
        "timings": [t.model_dump() for t in state.get("timings", [])],
        "sentences": {s.id: s.text for s in doc.sentences},
    }


def to_markdown(report: dict[str, Any]) -> str:
    p, d = report["process"], report["document"]
    out = [f"# Automation assessment — {p['name']}", "",
           f"{p['tasks']} tasks ({p['manual_tasks']} manual) across {len(p['lanes'])} lanes, {p['handoffs']} hand-offs. "
           f"Description: {d['sentences']} sentences; volume ≈ {d['facts']['monthly_volume'] or 'NI'}/month, "
           f"≈ {d['facts']['minutes_per_case'] or 'NI'} min/case.", "",
           f"**{len(report['candidates'])}** of {report['proposed']} proposed candidates survived critique and traceability.", "",
           "| # | Candidate | Technique | Impact | Feasibility | Evidence coverage | Quadrant |",
           "|---|---|---|---|---|---|---|"]
    for c in report["candidates"]:
        s = c["score"]
        out.append(f"| {c['id']} | {c['title'].split(': ', 1)[-1]} | {TECHNIQUE_LABEL[c['technique']]} | {s['impact']:.2f} | "
                   f"{s['feasibility']:.2f} | {s['coverage']:.0%} | {QUADRANT_LABEL[s['quadrant']]} |")
    for c in report["candidates"]:
        out += ["", f"## {c['id']} · {c['title']}", "", c["rationale"], "", "**Evidence**", ""]
        for e in c["evidence"]:
            out.append(f"- `{e['source']}:{e['ref']}` — “{e['quote']}”")
        out += ["", "| Question | Answer | Cites | Why |", "|---|---|---|---|"]
        for a in c["score"]["answers"]:
            out.append(f"| {a['question']} | {a['answer']} | {a['evidence_ref'] or '—'} | {a['justification']} |")
    if report["issues"]:
        out += ["", "## Review notes", ""]
        for i in report["issues"]:
            out.append(f"- **{i['severity']}** ({i['stage']}{', ' + i['candidate_id'] if i['candidate_id'] else ''}): {i['message']}")
    out += ["", "---", "_Scores are computed by a published rubric from cited answers; NI = not informed by the description._"]
    return "\n".join(out)


_CSS = """
:root{--bg:#000;--s:#070707;--l:#1c1c1c;--ink:#f4f4f0;--m:#7a7a76;--a:#c6ff00;--v:#a58bff;--w:#ffb020;--d:#ff4d5e}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.55 Inter,system-ui,sans-serif;padding:40px 16px}
main{max-width:1080px;margin:0 auto}h1{font:600 34px/1.1 'Space Grotesk',system-ui;letter-spacing:-.02em;margin:6px 0 10px}
h2{font:600 19px 'Space Grotesk',system-ui;margin:0}.label{font:500 10.5px ui-monospace,monospace;letter-spacing:.14em;text-transform:uppercase;color:var(--m)}
.card{border:1px solid var(--l);border-radius:3px;background:var(--s);margin:18px 0;padding:18px}
.chip{display:inline-block;border:1px solid #2a2a2a;border-radius:2px;padding:1px 7px;font:500 10.5px ui-monospace,monospace;text-transform:uppercase;letter-spacing:.06em}
.qw{color:var(--a);border-color:rgba(198,255,0,.4)}.st{color:var(--v);border-color:rgba(165,139,255,.4)}.fi{color:var(--w);border-color:rgba(255,176,32,.4)}.re{color:var(--d);border-color:rgba(255,77,94,.4)}
table{width:100%;border-collapse:collapse;margin-top:10px}td,th{text-align:left;padding:7px 9px;border-bottom:1px solid var(--l);vertical-align:top}
th{font:500 10.5px ui-monospace,monospace;letter-spacing:.12em;text-transform:uppercase;color:var(--m)}
.mono{font-family:ui-monospace,monospace;font-size:12.5px}.muted{color:var(--m)}q{color:var(--a)}
.matrix{position:relative;height:340px;border:1px solid var(--l);background:linear-gradient(var(--l),var(--l)) 55% 0/1px 100% no-repeat,linear-gradient(var(--l),var(--l)) 0 45%/100% 1px no-repeat}
.dot{position:absolute;transform:translate(-50%,50%);background:var(--a);color:#000;font:600 11px ui-monospace,monospace;padding:2px 6px;border-radius:2px}
.ax{position:absolute;font:10.5px ui-monospace,monospace;color:var(--m)}
.bar{height:4px;background:#121212;border-radius:2px;overflow:hidden;width:110px;display:inline-block;vertical-align:middle}.bar i{display:block;height:100%;background:var(--a)}
"""


def to_html(report: dict[str, Any]) -> str:
    e = html.escape
    p, d = report["process"], report["document"]
    qclass = {"quick_win": "qw", "strategic": "st", "fill_in": "fi", "reconsider": "re"}
    dots = "".join(
        f'<span class="dot" style="left:{c["score"]["feasibility"] * 100:.1f}%;bottom:{c["score"]["impact"] * 100:.1f}%" '
        f'title="{e(c["title"])}">{e(c["id"])}</span>' for c in report["candidates"])
    rows = "".join(
        f"<tr><td class=mono>{e(c['id'])}</td><td>{e(c['title'].split(': ', 1)[-1])}</td>"
        f"<td><span class=chip>{e(TECHNIQUE_LABEL[c['technique']])}</span></td>"
        f"<td><span class=bar><i style='width:{c['score']['impact'] * 100:.0f}%'></i></span> {c['score']['impact']:.2f}</td>"
        f"<td><span class=bar><i style='width:{c['score']['feasibility'] * 100:.0f}%'></i></span> {c['score']['feasibility']:.2f}</td>"
        f"<td class=mono>{c['score']['coverage']:.0%}</td>"
        f"<td><span class='chip {qclass[c['score']['quadrant']]}'>{QUADRANT_LABEL[c['score']['quadrant']]}</span></td></tr>"
        for c in report["candidates"])
    details = ""
    for c in report["candidates"]:
        ev = "".join(f"<li><span class=mono>{e(x['source'])}:{e(x['ref'])}</span> <q>{e(x['quote'])}</q></li>" for x in c["evidence"])
        qa = "".join(f"<tr><td class=mono>{e(a['question'])}</td><td><span class=chip>{e(a['answer'])}</span></td>"
                     f"<td class=mono>{e(a['evidence_ref'] or '—')}</td><td class=muted>{e(a['justification'])}</td></tr>"
                     for a in c["score"]["answers"])
        details += (f"<section class=card><div class=label>{e(c['id'])} · {e(c['technique'])}</div><h2>{e(c['title'])}</h2>"
                    f"<p class=muted>{e(c['rationale'])}</p><div class=label>evidence</div><ul>{ev}</ul>"
                    f"<table><tr><th>question</th><th>answer</th><th>cites</th><th>why</th></tr>{qa}</table></section>")
    issues = "".join(f"<tr><td><span class='chip {'re' if i['severity'] == 'blocker' else 'fi' if i['severity'] == 'warning' else ''}'>"
                     f"{e(i['severity'])}</span></td><td class=mono>{e(i['stage'])}</td><td class=mono>{e(i['candidate_id'] or '—')}</td>"
                     f"<td>{e(i['message'])}</td></tr>" for i in report["issues"])
    return f"""<!doctype html><html lang=en><head><meta charset=utf-8><meta name=viewport content="width=device-width,initial-scale=1">
<title>{e(p['name'])} — automation assessment</title><style>{_CSS}</style></head><body><main>
<div class=label>processlens · automation assessment</div><h1>{e(p['name'])}</h1>
<p class=muted>{p['tasks']} tasks ({p['manual_tasks']} manual) · {len(p['lanes'])} lanes · {p['handoffs']} hand-offs ·
≈ {d['facts']['monthly_volume'] or 'NI'} cases/month · ≈ {d['facts']['minutes_per_case'] or 'NI'} min/case ·
{len(report['candidates'])} of {report['proposed']} candidates survived traceability</p>
<section class=card><div class=label>priority matrix — impact × feasibility</div>
<div class=matrix>{dots}<span class=ax style="right:8px;bottom:6px">feasibility →</span><span class=ax style="left:8px;top:6px">↑ impact</span>
<span class=ax style="right:8px;top:6px;color:var(--a)">quick wins</span></div></section>
<section class=card><div class=label>candidates</div><table><tr><th>#</th><th>candidate</th><th>technique</th><th>impact</th>
<th>feasibility</th><th>evidence</th><th>quadrant</th></tr>{rows}</table></section>
{details}
<section class=card><div class=label>review notes (critic · traceability · audit)</div><table>{issues or '<tr><td class=muted>none</td></tr>'}</table></section>
<p class="muted mono">Scores come from a published rubric over cited answers. NI = not informed by the description.</p>
</main></body></html>"""


def to_docx(report: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)
    doc.add_heading(f"Automation assessment — {report['process']['name']}", 0)
    f = report["document"]["facts"]
    doc.add_paragraph(f"{report['process']['tasks']} tasks, {report['process']['handoffs']} hand-offs. "
                      f"Volume ≈ {f['monthly_volume'] or 'NI'}/month; ≈ {f['minutes_per_case'] or 'NI'} min/case.")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ["#", "Candidate", "Technique", "Impact", "Feasibility", "Quadrant"]):
        cell.text = text
    for c in report["candidates"]:
        s = c["score"]
        for cell, text in zip(table.add_row().cells, [c["id"], c["title"].split(": ", 1)[-1], TECHNIQUE_LABEL[c["technique"]],
                                                      f"{s['impact']:.2f}", f"{s['feasibility']:.2f}",
                                                      QUADRANT_LABEL[s["quadrant"]]]):
            cell.text = text
    for c in report["candidates"]:
        doc.add_heading(f"{c['id']} · {c['title']}", 2)
        doc.add_paragraph(c["rationale"])
        for ev in c["evidence"]:
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(f"{ev['source']}:{ev['ref']} ").bold = True
            para.add_run(f"“{ev['quote']}”")
        qa = doc.add_table(rows=1, cols=4)
        qa.style = "Light List"
        for cell, text in zip(qa.rows[0].cells, ["Question", "Answer", "Cites", "Why"]):
            cell.text = text
        for a in s["answers"] if (s := c["score"]) else []:
            for cell, text in zip(qa.add_row().cells, [a["question"], a["answer"], a["evidence_ref"] or "—", a["justification"]]):
                cell.text = text
    if report["issues"]:
        doc.add_heading("Review notes", 1)
        for i in report["issues"]:
            run = doc.add_paragraph(style="List Bullet").add_run(f"[{i['severity']}] {i['stage']}: {i['message']}")
            if i["severity"] == "blocker":
                run.font.color.rgb = RGBColor(0xC0, 0x1E, 0x2E)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def questions_reference() -> list[dict[str, Any]]:
    return [{k: q[k] for k in ("id", "axis", "weight", "text", "points")} for q in QUESTIONS]
